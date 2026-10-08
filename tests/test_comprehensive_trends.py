"""Synthetic contract fixtures only; never used as published evidence."""
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import pandas as pd
from openpyxl import load_workbook
from shoe_trends.ranking import prepare_observations, calculate_ranking, CHANNELS, DEFAULT_WEIGHTS
from shoe_trends.exporting import export_workbook, generate_upload_items, UPLOAD_COLUMNS

NOW = datetime(2026, 10, 8, 9, tzinfo=timezone.utc)
TOPICS = [{"id": "ballet", "label": "발레스니커즈", "pattern": "ballet|발레"}, {"id": "trail", "label": "트레일", "pattern": "trail|트레일"}]


def record(channel, text="ballet shoes", identity="one", **extra):
    return dict(source_type=channel, source_id=channel, source_url="https://example.com/" + identity,
                captured_at=NOW.isoformat(), published_at="2026-10-01T12:00:00+00:00", text=text,
                item_url="https://example.com/" + identity, rank=1, **extra)


def statuses(missing=()):
    return [dict(source_type=c, source_id=c, status="unavailable" if c in missing else "ok") for c in CHANNELS]


class ComprehensiveTrendTests(unittest.TestCase):
    def test_missing_is_not_zero_or_complete_score(self):
        obs, _ = prepare_observations([record("internal", keyword="발레스니커즈")], TOPICS, NOW)
        rank = calculate_ranking(obs, statuses(("sns", "newsletter", "ecommerce")))
        self.assertTrue(pd.isna(rank.iloc[0].SNS_Score))
        self.assertTrue(pd.isna(rank.iloc[0].Total_Trend_Score))
        self.assertEqual(rank.iloc[0].Evidence_Coverage, .2)
        self.assertAlmostEqual(rank.iloc[0].Provisional_Score, 100)

    def test_complete_four_channels_weighted_score_and_order_changes(self):
        rows = [record("internal", keyword="발레스니커즈"), record("ecommerce", "trail shoes"),
                record("sns", "trail shoes", likes=10), record("newsletter", "ballet shoes")]
        obs, _ = prepare_observations(rows, TOPICS, NOW)
        ranked = calculate_ranking(obs, statuses())
        self.assertEqual(ranked.iloc[0].Keyword, "트레일")
        self.assertEqual(ranked.iloc[0].Total_Trend_Score, 60)
        alternate = dict(internal=.4, ecommerce=.1, sns=.1, newsletter=.4)
        self.assertEqual(calculate_ranking(obs, statuses(), alternate).iloc[0].Keyword, "발레스니커즈")

    def test_duplicate_posts_across_hashtags_not_counted_twice(self):
        a = record("sns", post_id="same", likes=20)
        obs, raw = prepare_observations([a, {**a, "hashtag": "second"}], TOPICS, NOW)
        self.assertEqual(len(obs), 1)
        self.assertEqual(raw[1]["exclusion_reason"], "duplicate_entity")
        ranked = calculate_ranking(obs, statuses(("internal", "ecommerce", "newsletter")))
        self.assertEqual(ranked.iloc[0].SNS_Average_Likes, 20)

    def test_calendar_month_and_stale_captures_and_explicit_ranks(self):
        records = [record("newsletter"), record("newsletter", identity="old"), record("ecommerce", identity="rank51"), record("sns", identity="future"), record("internal", identity="stale", keyword="발레스니커즈")]
        records[1]["published_at"] = "2026-09-07T00:00:00+00:00"
        records[2]["rank"] = 51
        records[3]["published_at"] = "2026-10-09T00:00:00+00:00"
        records[4]["captured_at"] = "2026-09-01T00:00:00+00:00"
        obs, raw = prepare_observations(records, TOPICS, NOW)
        self.assertEqual(len(obs), 1)
        self.assertEqual(sum(r["included"] for r in raw), 1)

    def test_invalid_weights_rejected(self):
        obs, _ = prepare_observations([record("internal", keyword="발레스니커즈")], TOPICS, NOW)
        with self.assertRaises(ValueError):
            calculate_ranking(obs, statuses(), dict(internal=20, ecommerce=30, sns=30, newsletter=20))

    def test_partial_source_prevents_complete_claim(self):
        obs, _ = prepare_observations([record("internal", keyword="발레스니커즈")], TOPICS, NOW)
        states = statuses()
        states[1]["status"] = "partial"
        self.assertTrue(pd.isna(calculate_ranking(obs, states).iloc[0].Total_Trend_Score))

    def test_workbook_exact_schema_types_and_formula_injection(self):
        obs, raw = prepare_observations([record("internal", keyword="발레스니커즈")], TOPICS, NOW)
        ranked = calculate_ranking(obs, statuses(("sns",)))
        uploads = pd.DataFrame([["=CMD()", "Brand", "$120", "발레스니커즈", "https://example.com/shoe.jpg", "#123456", "@caption"]], columns=UPLOAD_COLUMNS)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "report.xlsx"
            export_workbook(ranked, uploads, raw, statuses(("sns",)), path, DEFAULT_WEIGHTS, NOW)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Combined_Ranking", "Upload_Format", "Raw_Data"])
            self.assertEqual(list(next(wb["Upload_Format"].values)), UPLOAD_COLUMNS)
            self.assertTrue(wb["Upload_Format"]["A1"].font.bold)
            self.assertEqual(wb["Upload_Format"]["A2"].data_type, "s")
            self.assertTrue(wb["Upload_Format"]["A2"].value.startswith("'="))
            self.assertIsNone(wb["Combined_Ranking"]["G2"].value)
            self.assertIsInstance(wb["Combined_Ranking"]["C2"].value, (float, int))
            wb.close()

    def test_upload_preserves_colors_and_does_not_invent_sns_description(self):
        obs, raw = prepare_observations([record("internal", keyword="발레스니커즈")], TOPICS, NOW)
        ranked = calculate_ranking(obs, statuses(("sns",)))
        shoe = dict(id="shoe", name="Ballet sneaker", brand="Test", style="TEST123", url="https://example.com/shoe", description="Official description", priceLabel="$100", colorSwatches=["#123456"], presentation={"image":"https://example.com/image.jpg"})
        model = dict(modelKey="test|ballet", name="Ballet sneaker", variants=[shoe, {**shoe, "id":"shoe-red"}], allVariantIds=["shoe", "shoe-red"])
        upload, items = generate_upload_items([model], ranked, TOPICS, raw)
        self.assertEqual(len(upload), 1)
        self.assertEqual(len(items[0]["variants"]), 2)
        self.assertEqual(items[0]["descriptionEvidence"]["type"], "official-catalog")
        self.assertFalse(items[0]["productPopularityVerified"])


if __name__ == "__main__":
    unittest.main()
